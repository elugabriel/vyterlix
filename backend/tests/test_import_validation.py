"""Checking an upload before importing: POST .../validate, GET .../validation, .../rows,
.../problems.csv."""

import csv as csv_module
import io
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import func, select

from app.core.config import Settings
from app.db.tenant import tenant_scope
from app.models.data import Customer, Product, Sale, Supplier
from app.models.identity import AuditLog, OrganizationUser, Role, User
from app.models.imports import DataImport, DataImportRow
from app.services import import_validation
from app.services import imports as imports_service

ORGS = "/api/v1/organizations"
D = Decimal
SALES_MAP = {
    "sold_on": "Date",
    "amount": "Total",
    "vat_amount": "VAT",
    "reference": "Order",
}
INC = {"vat_inclusive": True}


@pytest.fixture
def business(api, db, signup):
    auth = {
        "owner": signup("owner@acme.co.uk"),
        "viewer": signup("viewer@acme.co.uk"),
        "other": signup("owner@rival.co.uk"),
    }
    org_id = api.post(ORGS, json={"name": "Acme"}, headers=auth["owner"]).json()["id"]
    other_org = api.post(ORGS, json={"name": "Rival"}, headers=auth["other"]).json()["id"]
    db.add(
        OrganizationUser(
            organization_id=uuid.UUID(org_id),
            user_id=db.scalars(select(User.id).where(User.email == "viewer@acme.co.uk")).one(),
            role_id=db.scalars(select(Role.id).where(Role.code == "viewer")).first(),
        )
    )
    db.flush()
    return org_id, other_org, auth


def base(business, import_id, org=0):
    return f"{ORGS}/{business[org]}/imports/{import_id}"


def upload(api, business, content, name="till.csv", dataset="sales", org=0, who="owner"):
    res = api.post(
        f"{ORGS}/{business[org]}/imports",
        files={"file": (name, content)},
        data={"dataset": dataset},
        headers=business[2][who],
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


def map_it(api, business, import_id, mapping, options=INC, org=0, who="owner"):
    res = api.put(
        base(business, import_id, org) + "/mapping",
        json={"mapping": mapping, "options": options},
        headers=business[2][who],
    )
    assert res.status_code == 200, res.text


def validate(api, business, import_id, who="owner", org=0):
    return api.post(base(business, import_id, org) + "/validate", headers=business[2][who])


def mapped_upload(api, business, content, mapping=SALES_MAP, options=INC, **kw):
    import_id = upload(api, business, content, **kw)
    map_it(api, business, import_id, mapping, options, kw.get("org", 0), kw.get("who", "owner"))
    return import_id


def csv(*rows, header="Date,Order,Total,VAT"):
    return ("\n".join([header, *rows]) + "\n").encode()


GOOD_ROWS = (
    "28/09/2026,1001,12.00,2.00",
    "29/09/2026,1002,24.00,4.00",
    "30/09/2026,1003,6.00,1.00",
)


def rows_of(api, business, import_id, **params):
    res = api.get(base(business, import_id) + "/rows", params=params, headers=business[2]["owner"])
    assert res.status_code == 200, res.text
    return res.json()


def saved_sale(db, business, ref, source="csv", org=0):
    with tenant_scope(db, uuid.UUID(business[org])):
        db.add(
            Sale(
                sold_on=date(2026, 1, 1),
                net_amount=D("10"),
                vat_amount=D("2"),
                gross_amount=D("12"),
                source=source,
                source_ref=ref,
            )
        )
        db.flush()


# --- a clean file ----------------------------------------


def test_every_row_valid(api, db, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    res = validate(api, business, import_id)
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["rows"], body["valid"], body["invalid"], body["duplicate"]) == (3, 3, 0, 0)
    assert body["problems"] == [] and body["can_import"] is True
    assert body["totals"] == {"net": "35.00", "vat": "7.00", "gross": "42.00"}
    assert (body["date_from"], body["date_to"]) == ("2026-09-28", "2026-09-30")
    assert body["validated_at"]
    with tenant_scope(db, uuid.UUID(business[0])):
        imp = db.get(DataImport, uuid.UUID(import_id), populate_existing=True)
        assert imp.status == "validated"
        assert (imp.row_count, imp.valid_count, imp.invalid_count, imp.duplicate_count) == (
            3,
            3,
            0,
            0,
        )
        assert imp.imported_count == 0  # nothing is imported by checking
        assert db.scalar(select(func.count()).select_from(Sale)) == 0
        assert db.scalar(select(func.count()).select_from(DataImportRow)) == 3


def test_checking_is_audited(api, db, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    entry = db.scalars(select(AuditLog).where(AuditLog.action == "import.validated")).one()
    assert entry.target_id == import_id
    assert entry.details == {"rows": 3, "valid": 3, "invalid": 0, "duplicate": 0}


def test_the_summary_can_be_read_again(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    first = validate(api, business, import_id).json()
    again = api.get(base(business, import_id) + "/validation", headers=business[2]["owner"])
    assert again.status_code == 200 and again.json() == first


def test_warnings_for_data_gaps(api, business):
    mapping = {"sold_on": "Date", "amount": "Total", "vat_amount": "VAT"}
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS), mapping)
    warnings = {w["code"] for w in validate(api, business, import_id).json()["warnings"]}
    assert warnings == {"no_reference", "no_cost_of_goods"}


def test_a_file_with_a_reference_and_cost_has_no_warnings(api, business):
    mapping = SALES_MAP | {"cost": "Cost"}
    content = csv("28/09/2026,1001,12.00,2.00,3.00", header="Date,Order,Total,VAT,Cost")
    import_id = mapped_upload(api, business, content, mapping)
    assert validate(api, business, import_id).json()["warnings"] == []


# --- problems ----------------------------------------


MIXED = csv(
    "28/09/2026,1001,12.00,2.00",  # row 2: fine
    "31/04/2026,1002,12.00,2.00",  # row 3: not a real date
    "28/09/2026,1003,abc,2.00",  # row 4: unreadable amount
    "28/09/2026,1004,12.00,2.00",  # row 5: fine
    "29/09/2026,1005,£4,50,1.00",  # row 6: a comma for pence splits the cell: too many cells
    "28/09/2026,1006,12.00,2.00",  # row 7: fine
    "bad,1007,xx,2.00",  # row 8: two problems
)


def test_each_row_is_sorted_into_valid_or_invalid(api, business):
    import_id = mapped_upload(api, business, MIXED)
    body = validate(api, business, import_id).json()
    assert (body["rows"], body["valid"], body["invalid"], body["duplicate"]) == (7, 3, 4, 0)
    assert body["can_import"] is True
    kinds = {(p["code"], p["field"]): p for p in body["problems"]}
    assert kinds[("invalid_date", "sold_on")]["count"] == 2
    assert kinds[("invalid_date", "sold_on")]["rows"] == [3, 8]
    assert kinds[("invalid_amount", "amount")]["rows"] == [4, 8]
    assert "31/04/2026" in kinds[("invalid_date", "sold_on")]["example"]
    assert kinds[("invalid_date", "sold_on")]["example"].startswith("Row 3:")
    assert body["problems"] == sorted(body["problems"], key=lambda p: -p["count"])


def test_only_valid_rows_count_towards_the_totals(api, business):
    import_id = mapped_upload(api, business, MIXED)
    body = validate(api, business, import_id).json()
    assert body["totals"]["gross"] == "36.00"  # three valid rows of 12.00


def test_the_rows_can_be_listed_by_outcome(api, business):
    import_id = mapped_upload(api, business, MIXED)
    validate(api, business, import_id)
    bad = rows_of(api, business, import_id, status="invalid")
    assert bad["total"] == 4
    assert [r["row_number"] for r in bad["rows"]] == [3, 4, 6, 8]
    first = bad["rows"][0]
    assert first["status"] == "invalid"
    assert first["errors"][0]["field"] == "sold_on" and first["errors"][0]["code"] == "invalid_date"
    assert first["raw"] == {"Date": "31/04/2026", "Order": "1002", "Total": "12.00", "VAT": "2.00"}
    everything = rows_of(api, business, import_id)
    assert everything["total"] == 7 and len(everything["rows"]) == 7
    assert [r["status"] for r in everything["rows"]].count("valid") == 3


def test_listing_is_paged(api, business):
    import_id = mapped_upload(api, business, MIXED)
    validate(api, business, import_id)
    page = rows_of(api, business, import_id, limit=2, offset=2)
    assert page["total"] == 7 and [r["row_number"] for r in page["rows"]] == [4, 5]
    assert rows_of(api, business, import_id, offset=50)["rows"] == []


@pytest.mark.parametrize(
    "params", [{"status": "maybe"}, {"limit": 0}, {"limit": 201}, {"offset": -1}]
)
def test_bad_paging_parameters_are_refused(api, business, params):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    res = api.get(base(business, import_id) + "/rows", params=params, headers=business[2]["owner"])
    assert res.status_code == 422


def test_only_mapped_columns_are_kept_from_each_row(api, business):
    content = csv(
        "28/09/2026,1001,12.00,2.00,secret@example.com", header="Date,Order,Total,VAT,Phone"
    )
    import_id = mapped_upload(api, business, content)
    validate(api, business, import_id)
    raw = rows_of(api, business, import_id)["rows"][0]["raw"]
    assert set(raw) == {"Date", "Order", "Total", "VAT"}


def test_when_most_rows_have_problems_the_mapping_is_questioned(api, business):
    content = csv("28/09/2026,1,12,2", "nope,2,12,2", "nope,3,12,2", "nope,4,12,2")
    body = validate(api, business, mapped_upload(api, business, content)).json()
    assert "many_problems" in {w["code"] for w in body["warnings"]}


def test_a_file_where_nothing_can_be_imported(api, business):
    body = validate(api, business, mapped_upload(api, business, csv("nope,1,12,2"))).json()
    assert body["valid"] == 0 and body["can_import"] is False
    assert "nothing_to_import" in {w["code"] for w in body["warnings"]}
    assert body["totals"]["gross"] == "0.00" and body["date_from"] is None


# --- repeats ----------------------------------------


def test_an_identical_repeated_row_is_a_duplicate(api, business):
    content = csv(
        "28/09/2026,1001,12.00,2.00", "29/09/2026,1002,24.00,4.00", "28/09/2026,1001,12.00,2.00"
    )
    import_id = mapped_upload(api, business, content)
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["invalid"], body["duplicate"]) == (2, 0, 1)
    row = rows_of(api, business, import_id, status="duplicate")["rows"][0]
    assert row["row_number"] == 4
    assert row["errors"] == [
        {"field": "reference", "code": "duplicate_in_file", "message": "Same as row 2."}
    ]
    assert body["totals"]["gross"] == "36.00"  # the repeat is not counted twice


def test_the_same_reference_with_different_details_is_a_problem(api, business):
    content = csv("28/09/2026,1001,12.00,2.00", "28/09/2026,1001,99.00,2.00")
    import_id = mapped_upload(api, business, content)
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["invalid"], body["duplicate"]) == (1, 1, 0)
    assert body["problems"][0]["code"] == "repeated_reference"
    assert "row 2" in body["problems"][0]["example"]


def test_a_sale_already_in_the_data_is_a_duplicate(api, db, business):
    saved_sale(db, business, "1002")
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["duplicate"]) == (2, 1)
    [row] = rows_of(api, business, import_id, status="duplicate")["rows"]
    assert row["raw"]["Order"] == "1002" and row["errors"][0]["code"] == "already_imported"


def test_a_record_from_another_source_is_not_a_repeat(api, db, business):
    saved_sale(db, business, "1002", source="shopify")  # same number, different system
    body = validate(api, business, mapped_upload(api, business, csv(*GOOD_ROWS))).json()
    assert (body["valid"], body["duplicate"]) == (3, 0)


def test_another_businesss_record_is_not_a_repeat(api, db, business):
    saved_sale(db, business, "1002", org=1)
    body = validate(api, business, mapped_upload(api, business, csv(*GOOD_ROWS))).json()
    assert (body["valid"], body["duplicate"]) == (3, 0)


def test_rows_without_a_reference_are_never_called_repeats(api, business):
    mapping = {"sold_on": "Date", "amount": "Total", "vat_amount": "VAT"}
    content = csv("28/09/2026,1,12.00,2.00", "28/09/2026,1,12.00,2.00", "28/09/2026,1,12.00,2.00")
    body = validate(api, business, mapped_upload(api, business, content, mapping)).json()
    assert (body["valid"], body["duplicate"]) == (3, 0)  # three identical cash sales are fine


def test_blank_references_are_not_matched_with_each_other(api, business):
    content = csv("28/09/2026,,12.00,2.00", "28/09/2026,,12.00,2.00")
    body = validate(api, business, mapped_upload(api, business, content)).json()
    assert (body["valid"], body["duplicate"]) == (2, 0)


def test_repeats_are_found_across_the_batches_a_big_file_is_read_in(api, business, monkeypatch):
    monkeypatch.setattr(import_validation, "BATCH", 2)
    rows = [f"28/09/2026,{n},12.00,2.00" for n in range(1, 6)] + ["28/09/2026,1,12.00,2.00"]
    body = validate(api, business, mapped_upload(api, business, csv(*rows))).json()
    assert (body["rows"], body["valid"], body["duplicate"]) == (6, 5, 1)


def test_a_file_bigger_than_one_batch(api, business):
    rows = [f"28/09/2026,{n},12.00,2.00" for n in range(1, 2602)]
    body = validate(api, business, mapped_upload(api, business, csv(*rows))).json()
    assert (body["rows"], body["valid"]) == (2601, 2601)
    assert body["totals"]["gross"] == "31212.00"


# --- reading the file ----------------------------------------


def test_row_numbers_are_the_lines_in_the_file(api, business):
    content = (
        b"Date,Order,Total,VAT\n"
        b"28/09/2026,1001,12.00,2.00\n"
        b"\n"
        b'31/04/2026,"10\n02",12.00,2.00\n'
        b"28/09/2026,1003,nope,2.00\n"
    )
    import_id = mapped_upload(api, business, content)
    body = validate(api, business, import_id).json()
    rows = rows_of(api, business, import_id)["rows"]
    assert [r["row_number"] for r in rows] == [2, 4, 6]  # blank line skipped; cell spanning 4-5
    assert body["rows"] == 3


def test_headings_on_a_later_row(api, business):
    content = b"Till report\nRun 30/09/2026\nDate,Order,Total,VAT\n28/09/2026,1001,12.00,2.00\n"
    import_id = upload(api, business, content)
    api.patch(base(business, import_id), json={"header_row": 3}, headers=business[2]["owner"])
    map_it(api, business, import_id, SALES_MAP)
    body = validate(api, business, import_id).json()
    assert body["valid"] == 1
    assert rows_of(api, business, import_id)["rows"][0]["row_number"] == 4


def test_windows_1252_files_keep_their_pound_signs(api, business):
    content = "Date,Order,Total,VAT\n28/09/2026,1001,£12.00,£2.00\n".encode("cp1252")
    body = validate(api, business, mapped_upload(api, business, content)).json()
    assert body["valid"] == 1 and body["totals"]["gross"] == "12.00"


def test_excel_dates_and_numbers(api, business):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Date", "Order", "Total", "VAT"])
    sheet.append([datetime(2026, 9, 28), 1001, 12.0, 2.0])
    sheet.append([datetime(2026, 9, 29, 14, 30), 1002, 24.5, 4.08])
    sheet.append([None, None, None, None])
    sheet.append(["30/09/2026", "1003", "£6.00", "£1.00"])  # text cells in the same sheet
    out = io.BytesIO()
    book.save(out)
    import_id = mapped_upload(api, business, out.getvalue(), name="sales.xlsx")
    body = validate(api, business, import_id).json()
    assert (body["rows"], body["valid"], body["invalid"]) == (3, 3, 0)
    assert body["totals"]["gross"] == "42.50"
    assert [r["row_number"] for r in rows_of(api, business, import_id)["rows"]] == [2, 3, 5]


def test_the_chosen_excel_sheet_is_the_one_checked(api, business):
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    book.active.append(["Note"])
    book.active.append(["hi"])
    sheet = book.create_sheet("Sales")
    for row in (["Date", "Order", "Total", "VAT"], ["28/09/2026", 1, 12, 2]):
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    import_id = upload(api, business, out.getvalue(), name="b.xlsx")
    api.patch(base(business, import_id), json={"sheet_name": "Sales"}, headers=business[2]["owner"])
    map_it(api, business, import_id, SALES_MAP)
    assert validate(api, business, import_id).json()["valid"] == 1


# --- the other kinds of file ----------------------------------------


def test_expenses(api, business):
    content = b"Date,Supplier,Net,VAT\n03/09/2026,Flour Co,50.00,10.00\n04/09/2026,Gas,x,1\n"
    mapping = {"spent_on": "Date", "amount": "Net", "vat_amount": "VAT", "supplier": "Supplier"}
    import_id = mapped_upload(
        api, business, content, mapping, {"vat_inclusive": False}, dataset="expenses"
    )
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["invalid"]) == (1, 1)
    assert body["totals"] == {"net": "50.00", "vat": "10.00", "gross": "60.00"}


def test_customers_are_matched_to_existing_ones_by_email(api, db, business):
    with tenant_scope(db, uuid.UUID(business[0])):
        db.add(Customer(name="Jo", email="jo@example.com"))
        db.flush()
    content = b"Name,Email\nJo Bloggs,JO@example.com\nSam,sam@example.com\n,\nBad,nope\n"
    import_id = mapped_upload(
        api, business, content, {"name": "Name", "email": "Email"}, {}, dataset="customers"
    )
    body = validate(api, business, import_id).json()
    assert (body["rows"], body["valid"], body["duplicate"], body["invalid"]) == (3, 1, 1, 1)
    assert body["totals"] is None and body["date_from"] is None


def test_suppliers_are_matched_by_name_ignoring_case(api, db, business):
    with tenant_scope(db, uuid.UUID(business[0])):
        db.add(Supplier(name="Flour Co"))
        db.flush()
    content = b"Name\nFLOUR CO\nButter Ltd\nbutter ltd\n"
    import_id = mapped_upload(api, business, content, {"name": "Name"}, {}, dataset="suppliers")
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["duplicate"]) == (1, 2)


def test_products_are_matched_by_code(api, db, business):
    with tenant_scope(db, uuid.UUID(business[0])):
        db.add(Product(name="Loaf", sku="SD1"))
        db.flush()
    content = b"Name,SKU,Price\nSourdough,SD1,6.00\nBun,B1,1.20\nCake,,3.00\n"
    mapping = {"name": "Name", "sku": "SKU", "unit_price": "Price"}
    options = {"vat_inclusive": True, "default_vat_rate": "20"}
    import_id = mapped_upload(api, business, content, mapping, options, dataset="products")
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["duplicate"], body["invalid"]) == (2, 1, 0)


def test_stock_movements(api, business):
    content = (
        b"Date,SKU,Qty,Type\n"
        b"01/09/2026,SD1,40,opening\n"
        b"02/09/2026,SD1,-2,sale\n"
        b"03/09/2026,SD1,2,sale\n"
    )
    mapping = {"moved_on": "Date", "sku": "SKU", "quantity": "Qty", "movement_kind": "Type"}
    import_id = mapped_upload(api, business, content, mapping, {}, dataset="stock_movements")
    body = validate(api, business, import_id).json()
    assert (body["valid"], body["invalid"]) == (2, 1)
    assert body["problems"][0]["code"] == "sign_mismatch"


# --- running it again, and when it isn't allowed -------------------------


def test_checking_again_replaces_the_earlier_results(api, db, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    second = validate(api, business, import_id)
    assert second.status_code == 200 and second.json()["rows"] == 3
    with tenant_scope(db, uuid.UUID(business[0])):
        assert db.scalar(select(func.count()).select_from(DataImportRow)) == 3  # not 6


def test_a_new_mapping_discards_the_old_results(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    map_it(api, business, import_id, SALES_MAP, {"vat_inclusive": False})
    url = base(business, import_id)
    for suffix in ("/validation", "/rows", "/problems.csv"):
        res = api.get(url + suffix, headers=business[2]["owner"])
        assert res.status_code == 409 and res.json()["error"]["code"] == "not_validated"
    assert validate(api, business, import_id).json()["totals"]["gross"] == "49.00"  # now net + VAT


def test_changing_the_header_row_discards_the_old_results(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    api.patch(base(business, import_id), json={"header_row": 1}, headers=business[2]["owner"])
    res = api.get(base(business, import_id) + "/validation", headers=business[2]["owner"])
    assert res.status_code == 409


def test_the_columns_must_be_mapped_first(api, business):
    import_id = upload(api, business, csv(*GOOD_ROWS))
    res = validate(api, business, import_id)
    assert res.status_code == 409 and res.json()["error"]["code"] == "mapping_needed"


def test_an_import_that_has_run_cannot_be_checked_again(api, db, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    with tenant_scope(db, uuid.UUID(business[0])):
        imp = db.get(DataImport, uuid.UUID(import_id))
        imp.status, imp.imported_at = "imported", datetime.now(UTC)
        db.flush()
    res = validate(api, business, import_id)
    assert res.status_code == 409 and res.json()["error"]["code"] == "not_editable"


def test_the_original_file_must_still_exist(api, db, business, storage):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    with tenant_scope(db, uuid.UUID(business[0])):
        imp = db.get(DataImport, uuid.UUID(import_id))
        storage.delete(imp.storage_key)
        imp.file_deleted_at = datetime.now(UTC)
        db.flush()
    res = validate(api, business, import_id)
    assert res.status_code == 409 and res.json()["error"]["code"] == "file_deleted"


def test_a_failure_part_way_keeps_the_earlier_results(api, db, business, monkeypatch):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    before = validate(api, business, import_id).json()
    monkeypatch.setattr(
        imports_service,
        "get_settings",
        lambda: Settings(env="test", _env_file=None, max_upload_rows=1),
    )
    monkeypatch.setattr(import_validation, "get_settings", imports_service.get_settings)
    res = validate(api, business, import_id)
    assert res.status_code == 422 and res.json()["error"]["code"] == "too_many_rows"
    after = api.get(base(business, import_id) + "/validation", headers=business[2]["owner"])
    assert after.status_code == 200 and after.json() == before
    assert len(rows_of(api, business, import_id)["rows"]) == 3


def test_asking_for_results_before_checking(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    for suffix in ("/validation", "/rows", "/problems.csv"):
        res = api.get(base(business, import_id) + suffix, headers=business[2]["owner"])
        assert res.status_code == 409 and res.json()["error"]["code"] == "not_validated"


# --- the download of problems ----------------------------------------


def test_problems_download(api, business):
    import_id = mapped_upload(api, business, MIXED)
    validate(api, business, import_id)
    res = api.get(base(business, import_id) + "/problems.csv", headers=business[2]["owner"])
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert "attachment" in res.headers["content-disposition"]
    assert "problems" in res.headers["content-disposition"]
    text = res.content.decode("utf-8")
    assert text.startswith("﻿")  # so Excel reads it as UTF-8
    lines = text.lstrip("﻿").strip().splitlines()
    assert lines[0] == "Row,Status,Problems,Date,Total,VAT,Order"
    assert len(lines) == 1 + 4  # the four rows with problems, none of the valid ones
    assert lines[1].startswith("3,invalid,") and "31/04/2026" in lines[1]


def test_the_download_includes_repeats(api, business):
    content = csv("28/09/2026,1001,12.00,2.00", "28/09/2026,1001,12.00,2.00")
    import_id = mapped_upload(api, business, content)
    validate(api, business, import_id)
    res = api.get(base(business, import_id) + "/problems.csv", headers=business[2]["owner"])
    assert "3,duplicate,Same as row 2." in res.content.decode("utf-8")


def test_the_download_cannot_be_used_to_run_formulas_in_a_spreadsheet(api, business):
    content = csv(
        '=HYPERLINK("http://evil.example"),1001,12.00,2.00',
        "@SUM(1+1),1002,12.00,2.00",
        "+cmd|' /C calc'!A0,1003,12.00,2.00",
        "-5,1004,12.00,2.00",  # a plain negative number is left alone
    )
    import_id = mapped_upload(api, business, content)
    validate(api, business, import_id)
    res = api.get(base(business, import_id) + "/problems.csv", headers=business[2]["owner"])
    table = list(csv_module.reader(io.StringIO(res.content.decode("utf-8-sig"))))
    cells = [cell for row in table[1:] for cell in row]
    assert not any(cell.startswith(("=", "@", "+")) for cell in cells)
    assert any(cell.startswith("'=HYPERLINK") for cell in cells)
    assert any(cell.startswith("'@SUM") for cell in cells)
    assert any(cell.startswith("'+cmd") for cell in cells)
    assert "-5" in cells


def test_a_download_with_nothing_wrong_is_just_the_headings(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    res = api.get(base(business, import_id) + "/problems.csv", headers=business[2]["owner"])
    assert (
        res.content.decode("utf-8").strip().lstrip("﻿") == "Row,Status,Problems,Date,Total,VAT,Order"
    )


def test_the_download_name_is_safe(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS), name='till "x"\\y.csv')
    validate(api, business, import_id)
    res = api.get(base(business, import_id) + "/problems.csv", headers=business[2]["owner"])
    disposition = res.headers["content-disposition"]
    assert '"' not in disposition.split("filename*=")[1] and "\\" not in disposition


# --- who may do this ----------------------------------------


def test_viewers_cannot_check_or_read_results(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    assert validate(api, business, import_id, who="viewer").status_code == 403
    for suffix in ("/validation", "/rows", "/problems.csv"):
        res = api.get(base(business, import_id) + suffix, headers=business[2]["viewer"])
        assert res.status_code == 403


def test_logging_in_is_required(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    assert api.post(base(business, import_id) + "/validate").status_code == 401
    assert api.get(base(business, import_id) + "/rows").status_code == 401


def test_another_business_cannot_check_or_read_an_import(api, business):
    import_id = mapped_upload(api, business, csv(*GOOD_ROWS))
    validate(api, business, import_id)
    other = business[2]["other"]
    for org in (0, 1):
        assert validate(api, business, import_id, who="other", org=org).status_code == 404
        for suffix in ("/validation", "/rows", "/problems.csv"):
            res = api.get(base(business, import_id, org) + suffix, headers=other)
            assert res.status_code == 404
    assert validate(api, business, uuid.uuid4()).status_code == 404


def test_a_problem_lists_only_the_first_few_rows_but_counts_them_all(api, business):
    rows = [f"nope,{n},12.00,2.00" for n in range(1, 9)]
    body = validate(api, business, mapped_upload(api, business, csv(*rows))).json()
    [problem] = body["problems"]
    assert problem["count"] == 8 and problem["rows"] == [2, 3, 4, 5, 6]
